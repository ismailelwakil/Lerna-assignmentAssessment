"""Course + material service: secure upload, extraction (pypdf/docx/pptx/txt),
chunking with page/slide/section metadata, then indexing via the evidence
service (same Gemini embeddings + Qdrant instance as the existing projects)."""
from __future__ import annotations

import re
import uuid
from pathlib import Path

from ..config import get_settings
from ..db import db
from ..exceptions import ForbiddenError, NotFoundError, ValidationError
from ..models import AuditEvent, Course, CourseMaterial, MaterialChunk, User
from ..security import new_id, validate_material

CHUNK_SIZE = 1200
OVERLAP = 150


# ------------------------------------------------------------------ courses
def create_course(instructor: dict, title: str, subject: str = "",
                  code: str = "", description: str = "") -> dict:
    if instructor.get("role") != "instructor":
        raise ForbiddenError("Only instructors can create courses.")
    if not (title or "").strip():
        raise ValidationError("Course title is required.")
    with db().session_scope() as session:
        course = Course(id=new_id(), instructor_id=instructor["user_id"],
                        title=title.strip()[:200], subject=(subject or "")[:120],
                        code=(code or "")[:40], description=(description or "")[:2000])
        session.add(course)
        session.add(AuditEvent(actor_id=instructor["user_id"], event_type="course_created",
                               entity_type="course", entity_id=course.id,
                               detail={"title": course.title}))
        cid = course.id
    return {"course_id": cid, "title": course.title, "subject": course.subject}


def list_courses(instructor: dict) -> list[dict]:
    with db().session_scope() as session:
        query = session.query(Course)
        if instructor["role"] == "instructor":
            query = query.filter(Course.instructor_id == instructor["user_id"])
        courses = query.order_by(Course.created_at.desc()).all()
        result = []
        for course in courses:
            materials = session.query(CourseMaterial).filter_by(course_id=course.id).count()
            assessments = session.query(
                __import__("app.models", fromlist=["Assessment"]).Assessment).filter_by(
                course_id=course.id).count()
            result.append({"course_id": course.id, "title": course.title,
                           "subject": course.subject, "code": course.code,
                           "materials": materials, "assessments": assessments})
        return result


def owned_course(course_id: str, instructor: dict):
    with db().session_scope() as session:
        course = session.get(Course, course_id)
        if course is None:
            raise NotFoundError("Course not found.")
        if instructor["role"] == "instructor" and course.instructor_id != instructor["user_id"]:
            raise NotFoundError("Course not found.")  # uniform 404, no oracle
        return course


# -------------------------------------------------------------- extraction
def _chunk(text: str, *, page=None, slide=None, section=None) -> list[dict]:
    text = re.sub(r"[ \t]+", " ", text or "").strip()
    out, start = [], 0
    while start < len(text):
        piece = text[start:start + CHUNK_SIZE]
        if len(piece) == CHUNK_SIZE:
            cut = piece.rfind(" ")
            if cut > CHUNK_SIZE * 0.6:
                piece = piece[:cut]
        if piece.strip():
            out.append({"text": piece.strip(), "page": page, "slide": slide, "section": section})
        if start + CHUNK_SIZE >= len(text):
            break
        start += len(piece) - OVERLAP
    return out


def _extract(content: bytes, ext: str) -> tuple[list[dict], int]:
    import io
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        chunks = []
        for number, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:  # noqa: BLE001 — skip unreadable pages
                continue
            chunks += _chunk(text, page=number)
        return chunks, len(reader.pages)
    if ext == ".docx":
        from docx import Document
        document = Document(io.BytesIO(content))
        chunks, section, buffer = [], None, []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                if buffer:
                    chunks += _chunk(" ".join(buffer), section=section)
                    buffer = []
                section = text[:200]
            buffer.append(text)
        if buffer:
            chunks += _chunk(" ".join(buffer), section=section)
        return chunks, 0
    if ext == ".pptx":
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(content))
        chunks = []
        for index, slide in enumerate(presentation.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in paragraph.runs).strip()
                        if line:
                            parts.append(line)
            if parts:
                chunks += _chunk("\n".join(parts), slide=index)
        return chunks, len(presentation.slides)
    text = content.decode("utf-8", "replace")
    return _chunk(text), 1


def upload_material(instructor: dict, course_id: str, filename: str,
                    content: bytes) -> dict:
    course = owned_course(course_id, instructor)  # instructor-only ownership
    stem, ext = validate_material(filename, content)
    storage = Path(get_settings.__self__().db_url) if False else None  # noqa
    data_dir = Path(__file__).resolve().parent.parent.parent / "data" / "storage"
    data_dir.mkdir(parents=True, exist_ok=True)
    stored = data_dir / f"{uuid.uuid4().hex}{ext}"
    stored.write_bytes(content)

    with db().session_scope() as session:
        material = CourseMaterial(
            id=new_id(), course_id=course.id, filename=f"{stem}{ext}", ext=ext,
            size_bytes=len(content), storage_path=str(stored), status="uploaded")
        session.add(material)
        session.add(AuditEvent(actor_id=instructor["user_id"], event_type="material_uploaded",
                               entity_type="material", entity_id=material.id,
                               detail={"filename": material.filename, "bytes": len(content)}))
        mid = material.id
    return {"material_id": mid, "filename": f"{stem}{ext}", "status": "uploaded"}


def process_material(material_id: str) -> dict:
    """Extract → chunk → embed → Qdrant (evidence service). Synchronous for the
    prototype; status tracked for the UI."""
    from . import evidence_service
    with db().session_scope() as session:
        material = session.get(CourseMaterial, material_id)
        if material is None:
            raise NotFoundError("Material not found.")
        course_id = material.course_id
        content = Path(material.storage_path).read_bytes()
        material.status = "processing"

    try:
        chunks, pages = _extract(content, material_ext(session, material_id))
        if not chunks:
            raise ValidationError("No extractable text found in the file.")
        rows = []
        with db().session_scope() as session:
            session.query(MaterialChunk).filter_by(material_id=material_id).delete()
            material = session.get(CourseMaterial, material_id)
            for index, chunk in enumerate(chunks):
                rows.append(MaterialChunk(
                    id=new_id(), material_id=material_id, course_id=course_id,
                    chunk_index=index, text=chunk["text"], page=chunk["page"],
                    slide=chunk["slide"], section=chunk["section"]))
                session.add(rows[-1])
            material.status, material.page_count = "processed", pages
            material.chunk_count = len(rows)
            title = material.filename
        indexed = evidence_service.index_chunks(course_id, material_id, title, rows)
        with db().session_scope() as session:
            material = session.get(CourseMaterial, material_id)
            material.error = ""
        return {"material_id": material_id, "status": "processed",
                "chunks": len(rows), "pages": pages, "indexed": indexed}
    except Exception as exc:  # noqa: BLE001 — fail safely, keep state consistent
        with db().session_scope() as session:
            material = session.get(CourseMaterial, material_id)
            material.status, material.error = "failed", str(exc)[:300]
        raise


def material_ext(session, material_id: str) -> str:
    material = session.get(CourseMaterial, material_id)
    return material.ext if material else ""


def list_materials(course_id: str, instructor: dict) -> list[dict]:
    owned_course(course_id, instructor)
    with db().session_scope() as session:
        materials = session.query(CourseMaterial).filter_by(course_id=course_id).all()
        return [{"material_id": m.id, "filename": m.filename, "status": m.status,
                 "chunks": m.chunk_count, "pages": m.page_count,
                 "size_bytes": m.size_bytes, "error": m.error} for m in materials]
